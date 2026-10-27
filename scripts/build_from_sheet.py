#!/usr/bin/env python3
"""Regenerate data/players.json and data/results.json from the campaign Google Sheet export (.xlsx).

Usage:  python3 scripts/build_from_sheet.py path/to/sheet.xlsx
Needs:  pip install openpyxl

Privacy: the sheet's 'Discord UN' and 'Name' columns are NEVER written to the site.
Keep the .xlsx itself OUT of this folder (the site is public once hosted).
Faction columns: the sheet's Faction N slot position does not matter; each Faction/VP pair
is matched to a player by Lore Name, so a result typed into the "wrong" column still counts.
"""
import json, re, sys, datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
# Small spelling clean-ups applied to the public Army label (sheet left untouched).
ARMY_FIXES = {"Robot LEgions": "Robot Legions", "Dark Prime Brother’s": "Dark Prime Brothers",
              "High Elf fleets": "High Elf Fleets"}

def slug(s): return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
def num(v):
    if v is None or v == "": return None
    f = float(v); return int(f) if f.is_integer() else f

def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    pi = wb["Player Info"]; rows = list(pi.iter_rows(values_only=True))
    hdr = [str(h or "").strip() for h in rows[0]]
    players = []
    for r in rows[1:]:
        if not any(r): continue
        rec = dict(zip(hdr, r))
        lore = (rec.get("Lore Name") or "").strip()
        if not lore: continue
        army = (rec.get("Army") or "").strip()
        commander = (r[4] if len(r) > 4 else None) or rec.get("Commander") or ""
        players.append({"id": slug(lore), "loreName": lore, "army": ARMY_FIXES.get(army, army),
                        "commander": str(commander).strip()})
    by_lore = {p["loreName"].lower(): p for p in players}

    md = wb["Mission Data"]; rows = list(md.iter_rows(values_only=True))
    hdr = [str(h or "").strip() for h in rows[0]]
    missions, warnings = [], []
    for r in rows[1:]:
        if not any(v is not None for v in r): continue
        rec = dict(zip(hdr, r))
        d = rec.get("Date")
        date = d.date().isoformat() if isinstance(d, datetime.datetime) else str(d)
        res = []
        for i in range(1, 7):
            f, vp = rec.get(f"Faction {i}"), rec.get(f"Faction {i} VP")
            if not f: continue
            p = by_lore.get(str(f).strip().lower())
            if not p: warnings.append(f"Mission {rec['JAX OPR Mission']}: unknown faction '{f}'"); continue
            res.append({"player": p["id"], "vp": num(vp) or 0})
        n = num(rec.get("# Players"))
        if n and n != len(res):
            warnings.append(f"Mission {num(rec['JAX OPR Mission'])}: '# Players' says {n} but {len(res)} faction results found")
        res.sort(key=lambda x: -x["vp"])
        missions.append({"mission": num(rec["JAX OPR Mission"]), "date": date,
                         "voidsunMission": num(rec.get("Voidsun Mission")), "playerCount": n, "results": res})
    missions.sort(key=lambda m: m["mission"])
    (DATA / "players.json").write_text(json.dumps(players, indent=2, ensure_ascii=False) + "\n")
    (DATA / "results.json").write_text(json.dumps({"generated": datetime.date.today().isoformat(),
        "missions": missions}, indent=2, ensure_ascii=False) + "\n")
    print(f"Wrote {len(players)} players, {len(missions)} mission results.")
    for w in warnings: print("WARNING:", w)

if __name__ == "__main__":
    if len(sys.argv) != 2: sys.exit(__doc__)
    main(sys.argv[1])
